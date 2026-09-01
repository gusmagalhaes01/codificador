import datetime
import os
import re
import shutil
import sys
import tempfile
import unittest

_RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if _RAIZ not in sys.path:
    sys.path.insert(0, _RAIZ)
_AQUI = os.path.dirname(os.path.abspath(__file__))
if _AQUI not in sys.path:
    sys.path.insert(0, _AQUI)

import logica as app
from cadastro_teste import CADASTRO_TESTE

KLOSTERS = "01195716000154"


def ler(nome):
    caminho = os.path.join(_AQUI, "dados", nome)
    with open(caminho, encoding="utf-8") as f:
        return f.read()


class TestExtracaoNfse(unittest.TestCase):
    """Campos lidos da NFS-e da F&F (fixture nfse_ff.txt, tomador KLOSTERS)."""

    def setUp(self):
        self.dados = app.extrair_dados_nfse(ler("nfse_ff.txt"))

    def test_identificacao_da_nota(self):
        self.assertEqual(self.dados["numero"], "7884")
        self.assertEqual(self.dados["competencia"], datetime.date(2026, 6, 25))
        self.assertEqual(self.dados["emissao"], datetime.datetime(2026, 6, 25, 0, 32, 31))

    def test_tomador_nao_pega_dados_do_emitente(self):
        # O bloco do EMITENTE vem antes e tem os mesmos rótulos (CNPJ,
        # "Nome / Nome Empresarial"). Sem o recorte por seção, a leitura
        # devolveria a F&F em vez do condomínio.
        self.assertEqual(self.dados["cnpj_tomador"], KLOSTERS)
        self.assertEqual(self.dados["nome_tomador"], "KLOSTERS")

    def test_valores(self):
        self.assertEqual(self.dados["valor_servico"], 78.61)
        self.assertEqual(self.dados["valor_liquido"], 78.61)
        self.assertEqual(self.dados["bc_issqn"], 78.61)
        self.assertEqual(self.dados["aliquota"], 5.0)
        self.assertEqual(self.dados["issqn"], 3.93)
        self.assertEqual(self.dados["retencao_issqn"], "Não Retido")

    def test_sem_retencao_federal_vira_vazio(self):
        # Nota sem retenção: os dois campos vêm como "-" no DANFSe. Têm de
        # virar célula vazia, nunca 0 — 0 significaria "reteve zero".
        self.assertIsNone(self.dados["previdencia_retida"])
        self.assertIsNone(self.dados["contrib_sociais_retidas"])
        self.assertEqual(self.dados["valor_liquido"], self.dados["valor_servico"])

    def test_issqn_bate_com_bc_vezes_aliquota(self):
        # Conferência independente: se algum dos três fosse lido errado,
        # a conta não fecharia.
        calculado = round(self.dados["bc_issqn"] * self.dados["aliquota"] / 100, 2)
        self.assertAlmostEqual(calculado, self.dados["issqn"], places=2)

    def test_documento_que_nao_e_nfse(self):
        # Boleto não tem "Número da NFS-e" — deve ser recusado, não chutado.
        self.assertIsNone(app.extrair_dados_nfse(ler("boleto_avulso.txt")))


class TestDanfseV2(unittest.TestCase):
    """
    A Prefeitura passou a emitir DANFSe v2.0 (fixture nfse_v2_ibs.txt, NF real
    do lote "beneficios nf") convivendo no mesmo dia com notas v1.0 — não foi
    uma troca limpa. Diferenças que quebravam a extração:
    - Rótulos de identificação (Número/Competência/Data de emissão da NFS-e)
      e os títulos de seção saem em CAIXA ALTA no v2.0 ("NÚMERO DA NFS-E"),
      Título Normal no v1.0 ("Número da NFS-e") — campo_danfse/bloco_secao
      eram sensíveis a maiúsculas e não achavam nada, por isso a extração
      inteira desistia logo no "numero".
    - A seção do tomador mudou de nome: "TOMADOR DO SERVIÇO" (v1.0) virou
      "TOMADOR / ADQUIRENTE" (v2.0) — não é só maiúscula, é texto diferente.
    - O campo do valor do serviço na seção de totais mudou de "Valor do
      Serviço" para "Valor da Operação / Serviço".
    IBS/CBS (tributos novos da reforma tributária, só existem no v2.0) não
    são extraídos de propósito — não pedidos ainda.
    """

    def setUp(self):
        self.dados = app.extrair_dados_nfse(ler("nfse_v2_ibs.txt"))

    def test_nota_v2_e_reconhecida(self):
        self.assertIsNotNone(self.dados)

    def test_identificacao_da_nota(self):
        self.assertEqual(self.dados["numero"], "95631")
        self.assertEqual(self.dados["competencia"], datetime.date(2026, 7, 21))
        self.assertEqual(self.dados["emissao"], datetime.datetime(2026, 7, 21, 10, 2, 24))

    def test_tomador(self):
        self.assertEqual(self.dados["cnpj_tomador"], "68584267000107")
        self.assertEqual(self.dados["nome_tomador"], "CONDOMINIO DO EDIFICIO ULYSSEA")

    def test_valor_do_servico_com_rotulo_novo(self):
        self.assertEqual(self.dados["valor_servico"], 133.32)
        self.assertEqual(self.dados["valor_liquido"], 133.32)

    def test_issqn(self):
        self.assertEqual(self.dados["bc_issqn"], 133.32)
        self.assertEqual(self.dados["aliquota"], 5.0)
        self.assertEqual(self.dados["issqn"], 6.67)


class TestRotuloQuebradoNoHifen(unittest.TestCase):
    """
    DANFSe v2.0 em que "NFS-e" QUEBRA A LINHA entre o hífen e o "e" (fixture
    nfse_v2_rotulo_quebrado.txt, ESOCIAL 10002 San Remo, agosto/2026). O
    rótulo cai na borda da coluna e sai do extrator como dois blocos em
    linhas diferentes:

        NÚMERO DA NFS-
        e

        12815

    As notas de julho, mesma v2.0, traziam "NÚMERO DA NFS-E" numa linha só —
    por isso os 668/668 da validação da v6.11.0 não pegaram este caso.

    O estrago era desproporcional: `campo_danfse` exigia o rótulo inteiro
    numa linha, então "numero" vinha None e `extrair_dados_nfse` devolvia
    None logo na primeira checagem. A nota inteira era descartada como "não é
    DANFSe" — o MESMO desfecho de um "Detalhamento do Faturamento", que é
    documento que de fato não deve ser lido. Nada na tela distinguia os dois.
    """

    def setUp(self):
        self.dados = app.extrair_dados_nfse(ler("nfse_v2_rotulo_quebrado.txt"))

    def test_nota_com_rotulo_quebrado_e_reconhecida(self):
        # A regressão que motivou o fix: antes dele, isto era None.
        self.assertIsNotNone(self.dados)

    def test_identificacao_atravessa_a_quebra(self):
        # Os três rótulos afetados são justamente os do cabeçalho — os
        # únicos que terminam em "NFS-e" e caem na borda da coluna.
        self.assertEqual(self.dados["numero"], "12815")
        self.assertEqual(self.dados["competencia"], datetime.date(2026, 8, 24))
        self.assertEqual(self.dados["emissao"], datetime.datetime(2026, 8, 24, 23, 5, 8))

    def test_tomador(self):
        self.assertEqual(self.dados["cnpj_tomador"], "08578541000103")
        self.assertEqual(self.dados["nome_tomador"], "SAN REMO")

    def test_valores(self):
        self.assertEqual(self.dados["valor_servico"], 21.31)
        self.assertEqual(self.dados["valor_liquido"], 21.31)
        self.assertEqual(self.dados["bc_issqn"], 21.31)
        self.assertEqual(self.dados["aliquota"], 5.0)
        self.assertEqual(self.dados["issqn"], 1.07)


class TestPadraoRotulo(unittest.TestCase):
    """
    `padrao_rotulo` tolera UMA quebra dentro de "NFS-e", nunca mais que isso.
    A folga tinha que ser estreita: com `\\s*` no lugar de `\\n?`, um rótulo
    seguido de campo vazio casaria com a palavra iniciada em "e" da linha
    seguinte (ex: "emissão") e devolveria o valor do campo errado — o mesmo
    risco que já fez `campo_danfse` recusar `\\s*` solto.
    """

    def test_rotulo_sem_nfse_fica_intacto(self):
        self.assertEqual(app.padrao_rotulo("BC ISSQN"), re.escape("BC ISSQN"))

    def test_le_rotulo_quebrado_e_inteiro_do_mesmo_jeito(self):
        for bloco in ("Número da NFS-\ne\n \n12815", "NÚMERO DA NFS-E\n \n12815"):
            with self.subTest(bloco=bloco):
                self.assertEqual(app.campo_danfse(bloco, "Número da NFS-e"), "12815")

    def test_nao_atravessa_duas_quebras(self):
        # Campo vazio: o valor NÃO é o rótulo de baixo. Sem o limite de uma
        # quebra, "emissão..." seria engolido como se fosse o "e" do rótulo.
        bloco = "Número da NFS-\n\n\nemissão da nota\n \n24/08/2026"
        self.assertIsNone(app.campo_danfse(bloco, "Número da NFS-e"))


class TestRetencaoFederal(unittest.TestCase):
    """
    Nota com retenção ("3 - PIS/COFINS/CSLL Retidos", fixture nfse_ff_retido).
    "Contribuições Sociais - Retidas" já é o agregado de PIS+COFINS+CSLL — é o
    que desconta da nota. Os campos "PIS/COFINS - Débito Apuração Própria" do
    DANFSe são débito próprio da empresa e NÃO entram na planilha, justamente
    para não serem confundidos com retenção.
    """

    def setUp(self):
        self.dados = app.extrair_dados_nfse(ler("nfse_ff_retido.txt"))

    def test_contribuicoes_sociais_retidas(self):
        self.assertEqual(self.dados["contrib_sociais_retidas"], 15.10)

    def test_previdenciaria_nao_retida_fica_vazia(self):
        self.assertIsNone(self.dados["previdencia_retida"])

    def test_retencao_explica_a_diferenca_ate_o_liquido(self):
        # Conferência independente: serviço − retenções = líquido.
        # Se a retenção fosse lida do campo errado (ex: o PIS de apuração
        # própria, R$ 2,11), esta conta não fecharia.
        retido = (self.dados["contrib_sociais_retidas"] or 0) + \
                 (self.dados["previdencia_retida"] or 0)
        self.assertAlmostEqual(
            self.dados["valor_servico"] - retido, self.dados["valor_liquido"], places=2)

    def test_nao_confunde_com_pis_cofins_de_apuracao_propria(self):
        # PIS 2,11 e COFINS 9,74 existem no documento; nenhum dos dois pode
        # aparecer como retenção.
        self.assertNotIn(self.dados["contrib_sociais_retidas"], (2.11, 9.74, 11.85))


class TestCampoVazioNaoPuxaRotuloSeguinte(unittest.TestCase):
    """
    "Benefício Municipal" vem sem valor no DANFSe, seguido direto pelo rótulo
    "Valor do Serviço". Um regex frouxo devolveria "Valor do Serviço" como se
    fosse o valor do benefício — e num campo monetário isso viraria lixo na
    planilha. A conversão é a última linha de defesa.
    """

    def test_conversao_rejeita_rotulo_lido_como_valor(self):
        bloco = app.bloco_secao(ler("nfse_ff.txt"), "TRIBUTAÇÃO MUNICIPAL")
        bruto = app.campo_danfse(bloco, "Benefício Municipal")
        self.assertIsNone(app.converter_valor_br(bruto))

    def test_valor_do_servico_vem_da_secao_de_totais(self):
        texto = ler("nfse_ff.txt")
        bloco_total = app.bloco_secao(texto, "VALOR TOTAL DA NFS")
        self.assertNotIn("TRIBUTAÇÃO", bloco_total)
        self.assertEqual(
            app.converter_valor_br(app.campo_danfse(bloco_total, "Valor do Serviço")), 78.61)


class TestConversores(unittest.TestCase):
    def test_valor_br(self):
        self.assertEqual(app.converter_valor_br("R$ 1.234,56"), 1234.56)
        self.assertEqual(app.converter_valor_br("R$ 0,00"), 0.0)
        self.assertIsNone(app.converter_valor_br("-"))
        self.assertIsNone(app.converter_valor_br(""))
        self.assertIsNone(app.converter_valor_br(None))
        self.assertIsNone(app.converter_valor_br("Não Retido"))

    def test_percentual(self):
        self.assertEqual(app.converter_percentual("5,00 %"), 5.0)
        self.assertEqual(app.converter_percentual("2,5%"), 2.5)
        self.assertIsNone(app.converter_percentual("-"))

    def test_data(self):
        self.assertEqual(app.converter_data_br("21/07/2026"), datetime.date(2026, 7, 21))
        self.assertEqual(app.converter_data_br("21/07/2026 20:35:22"),
                         datetime.datetime(2026, 7, 21, 20, 35, 22))
        self.assertIsNone(app.converter_data_br("-"))
        self.assertIsNone(app.converter_data_br("32/13/2026"))


class TestLinhaDaPlanilha(unittest.TestCase):
    def setUp(self):
        self.dados = app.extrair_dados_nfse(ler("nfse_ff.txt"))

    def test_codigo_vem_do_cadastro_pelo_cnpj(self):
        linha = app.linha_planilha_nfse("PGR 10004 Klosters.pdf", self.dados, CADASTRO_TESTE)
        self.assertEqual(linha[0], "PGR 10004 Klosters.pdf")
        self.assertEqual(linha[6], "10004")     # código do cadastro
        self.assertEqual(linha[15], "")         # sem observação

    def test_cnpj_fora_do_cadastro_avisa_e_nao_inventa_codigo(self):
        linha = app.linha_planilha_nfse("nota.pdf", self.dados, {})
        self.assertEqual(linha[6], "")
        self.assertIn("não está no cadastro", linha[15])

    def test_documento_ignorado_vira_linha_so_com_motivo(self):
        linha = app.linha_planilha_nfse("outro.pdf", None, CADASTRO_TESTE, "Não é uma NFS-e")
        self.assertEqual(linha[0], "outro.pdf")
        self.assertEqual(linha[15], "Não é uma NFS-e")
        self.assertTrue(all(c is None for c in linha[1:15]))


class TestPlanilhaGerada(unittest.TestCase):
    def setUp(self):
        self.pasta = tempfile.mkdtemp(prefix="teste_nfse_")
        self.destino = os.path.join(self.pasta, "saida.xlsx")

    def tearDown(self):
        shutil.rmtree(self.pasta, ignore_errors=True)

    def test_valores_e_datas_sao_numeros_e_datas_nao_texto(self):
        from openpyxl import load_workbook

        dados = app.extrair_dados_nfse(ler("nfse_ff.txt"))
        linha = app.linha_planilha_nfse("PGR 10004 Klosters.pdf", dados, CADASTRO_TESTE)
        app.salvar_planilha_nfse(self.destino, [linha])

        sheet = load_workbook(self.destino).active
        self.assertEqual(sheet.max_row, 2)
        self.assertEqual([c.value for c in sheet[1]], [c[0] for c in app.COLUNAS_NFSE])

        gravada = [c.value for c in sheet[2]]
        # Somar no Excel só funciona se for número de verdade, não texto
        self.assertIsInstance(gravada[7], (int, float))
        self.assertEqual(gravada[7], 78.61)
        self.assertIsInstance(gravada[2], datetime.datetime)
        self.assertEqual(sheet.freeze_panes, "A2")
        self.assertIsNotNone(sheet.auto_filter.ref)


class TestTomadorPessoaFisica(unittest.TestCase):
    """Condomínio sem CNPJ próprio aparece na nota com o CPF do síndico no
    bloco do tomador. Sem ler os dois formatos, a planilha de notas desses
    condomínios sairia sistematicamente sem código."""

    def setUp(self):
        #  Mesma nota da fixture, com o documento do tomador trocado por um
        #  CPF. O rótulo do DANFSe já é "CNPJ / CPF / NIF" — só o valor muda.
        self.texto = ler("nfse_ff.txt").replace("01.195.716/0001-54",
                                                 "529.982.247-25")

    def test_acha_cpf_no_bloco_do_tomador(self):
        dados = app.extrair_dados_nfse(self.texto)
        self.assertIsNotNone(dados)
        self.assertEqual(dados["cnpj_tomador"], "52998224725")

    def test_cpf_com_dv_errado_nao_entra(self):
        texto = self.texto.replace("529.982.247-25", "529.982.247-24")
        dados = app.extrair_dados_nfse(texto)
        self.assertIsNotNone(dados)
        self.assertEqual(dados["cnpj_tomador"], "")

    def test_nota_com_cnpj_continua_lendo_o_cnpj(self):
        """A fixture original não pode mudar de resultado."""
        dados = app.extrair_dados_nfse(ler("nfse_ff.txt"))
        self.assertEqual(dados["cnpj_tomador"], KLOSTERS)


if __name__ == "__main__":
    unittest.main()
