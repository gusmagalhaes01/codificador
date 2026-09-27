# -*- coding: utf-8 -*-
"""Lançamento no Superlógica direto da extração (aba 2)."""
import datetime
import os
import sys
import unittest

_RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if _RAIZ not in sys.path:
    sys.path.insert(0, _RAIZ)
_AQUI = os.path.dirname(os.path.abspath(__file__))
if _AQUI not in sys.path:
    sys.path.insert(0, _AQUI)

import logica as app

#  Texto nativo como o pypdf devolve um DARF do DCTFWeb (sintético, mesmo
#  layout das guias reais). O vencimento aparece três vezes: no cabeçalho,
#  na composição e no recibo do rodapé.
DARF_OK = (
    "Documento de Arrecadação\nde Receitas Federais\n"
    "Período de Apuração Data de Vencimento Número do Documento\n"
    "07.16.26261.0000000-0 Pagar este documento até\n"
    "18/09/2026Observações\n"
    "Valor Total do Documento\n58,51\n"
    "5952 RET DE CONTRIBUICOES PAGT PJ A PJ DE DIR PRIV 58,51 58,51\n"
    "PA:08/2026 Vencimento:18/09/2026\n"
    "SENDA (Versão:1.5.10) 18/09/2026 09:46:491 1Página: /\n"
    "Número: 07.16.26261.0000000-0\n"
    "Pagar até: 18/09/2026\n"
    "Valor: 58,51\n"
)


class TestVencimentoDoDarf(unittest.TestCase):

    def test_tres_ocorrencias_iguais_dao_a_data(self):
        self.assertEqual(app.vencimento_do_darf(DARF_OK), datetime.date(2026, 9, 18))

    def test_ocorrencias_divergentes_nao_escolhem_nenhuma(self):
        texto = DARF_OK.replace("Pagar até: 18/09/2026", "Pagar até: 19/09/2026")
        self.assertIsNone(app.vencimento_do_darf(texto))

    def test_uma_ocorrencia_so_nao_basta(self):
        #  A conferência existe porque a data se repete; com uma só, não há
        #  com o que conferir.
        texto = ("Pagar até: 18/09/2026\n"
                 "5952 RET DE CONTRIBUICOES 58,51\n")
        self.assertIsNone(app.vencimento_do_darf(texto))

    def test_nenhuma_ocorrencia(self):
        self.assertIsNone(app.vencimento_do_darf("Documento de Arrecadação sem data"))

    def test_data_que_nao_existe_e_recusada(self):
        texto = DARF_OK.replace("18/09/2026", "31/02/2026")
        self.assertIsNone(app.vencimento_do_darf(texto))

    def test_carimbo_de_emissao_nao_conta(self):
        #  "SENDA ... 18/09/2026 09:46" é quando a guia foi emitida, sem
        #  rótulo de vencimento. Se contasse, trocar a data de emissão
        #  mudaria o resultado.
        texto = DARF_OK.replace("SENDA (Versão:1.5.10) 18/09/2026",
                                "SENDA (Versão:1.5.10) 01/01/2026")
        self.assertEqual(app.vencimento_do_darf(texto), datetime.date(2026, 9, 18))

    def test_multa_apos_vencimento_do_boleto_nao_casa(self):
        texto = "APOS VENCIMENTO MULTA DE 2% e JUROS DE 1%\n10/09/2026\n"
        self.assertIsNone(app.vencimento_do_darf(texto))

    def test_texto_vazio_ou_none(self):
        self.assertIsNone(app.vencimento_do_darf(""))
        self.assertIsNone(app.vencimento_do_darf(None))


from cadastro_teste import CADASTRO_TESTE as CADASTRO

KLOSTERS = "01195716000154"      # código 10004, ID SL 44
LAGO = "07945453000130"          # código 10590, SEM ID SL

NFSE_SEM_RETENCAO = {
    "numero": "12367", "competencia": datetime.date(2026, 8, 24),
    "cnpj_tomador": KLOSTERS, "valor_servico": 21.31, "valor_liquido": 21.31,
    "previdencia_retida": None, "contrib_sociais_retidas": None,
}
BOLETO = {
    "tipo": "Boleto bancário", "valor": 82.9,
    "vencimento": datetime.date(2026, 9, 10),
    "linha_digitavel": "23790472089000015979956006290706515650000008290",
    "documento_pagador": KLOSTERS,
}
ARRECADACAO = {
    "tipo": "Arrecadação", "valor": 58.51, "vencimento": None,
    "linha_digitavel": "858300000009585103852620610716262610830654448534",
    "documento_pagador": KLOSTERS,
}


class TestDocumentoParaLancamento(unittest.TestCase):

    def test_nota_fiscal(self):
        d = app.documento_para_lancamento("n.pdf", dados_nfse=NFSE_SEM_RETENCAO)
        self.assertEqual(d["tipo"], app.TIPO_LANCAMENTO_NOTA)
        self.assertEqual(d["documento"], KLOSTERS)
        self.assertEqual(d["valor"], 21.31)
        self.assertIsNone(d["vencimento"])          # vem do lote, não da nota
        self.assertIsNone(d["linha_digitavel"])
        self.assertEqual(d["numero_documento"], "12367")
        self.assertEqual(d["competencia"], datetime.date(2026, 8, 24))
        self.assertFalse(d["retencao"])

    def test_retencao_pela_diferenca_entre_servico_e_liquido(self):
        dados = dict(NFSE_SEM_RETENCAO, valor_servico=217.33, valor_liquido=207.22)
        self.assertTrue(app.documento_para_lancamento("n.pdf", dados_nfse=dados)["retencao"])

    def test_retencao_pelas_contribuicoes_sociais_retidas(self):
        #  O segundo sinal sozinho: serviço e líquido iguais, mas o campo de
        #  retenção preenchido. Basta um dos dois.
        dados = dict(NFSE_SEM_RETENCAO, contrib_sociais_retidas=10.11)
        self.assertTrue(app.documento_para_lancamento("n.pdf", dados_nfse=dados)["retencao"])

    def test_retencao_pela_previdencia_retida(self):
        dados = dict(NFSE_SEM_RETENCAO, previdencia_retida=2.34)
        self.assertTrue(app.documento_para_lancamento("n.pdf", dados_nfse=dados)["retencao"])

    def test_boleto(self):
        d = app.documento_para_lancamento("b.pdf", dados_boleto=BOLETO)
        self.assertEqual(d["tipo"], app.TIPO_LANCAMENTO_BOLETO)
        self.assertEqual(d["vencimento"], datetime.date(2026, 9, 10))
        self.assertEqual(d["linha_digitavel"], BOLETO["linha_digitavel"])
        self.assertIsNone(d["numero_documento"])
        self.assertIsNone(d["competencia"])

    def test_arrecadacao_pega_o_vencimento_do_texto(self):
        d = app.documento_para_lancamento("g.pdf", dados_boleto=ARRECADACAO, texto=DARF_OK)
        self.assertEqual(d["tipo"], app.TIPO_LANCAMENTO_ARRECADACAO)
        self.assertEqual(d["vencimento"], datetime.date(2026, 9, 18))
        self.assertIsNone(d["numero_documento"])
        self.assertIsNone(d["competencia"])

    def test_documento_nao_reconhecido(self):
        self.assertIsNone(app.documento_para_lancamento("x.pdf"))


class TestMotivoForaDoLancamento(unittest.TestCase):

    def doc(self, **mudancas):
        base = app.documento_para_lancamento("b.pdf", dados_boleto=BOLETO)
        base.update(mudancas)
        return base

    def test_documento_completo_pode_ser_lancado(self):
        self.assertEqual(app.motivo_fora_do_lancamento(self.doc(), CADASTRO), "")

    def test_nota_com_retencao(self):
        d = app.documento_para_lancamento(
            "n.pdf", dados_nfse=dict(NFSE_SEM_RETENCAO, contrib_sociais_retidas=10.11))
        self.assertIn("retenção", app.motivo_fora_do_lancamento(d, CADASTRO))

    def test_condominio_nao_identificado(self):
        self.assertIn("não identificado", app.motivo_fora_do_lancamento(
            self.doc(documento=None), CADASTRO))

    def test_documento_fora_do_cadastro(self):
        self.assertIn("não identificado", app.motivo_fora_do_lancamento(
            self.doc(documento="11222333000181"), CADASTRO))

    def test_condominio_sem_id_sl(self):
        self.assertIn("ID SL", app.motivo_fora_do_lancamento(
            self.doc(documento=LAGO), CADASTRO))

    def test_sem_valor(self):
        self.assertIn("valor", app.motivo_fora_do_lancamento(
            self.doc(valor=None), CADASTRO))

    def test_boleto_sem_vencimento(self):
        #  Fator 0000/9999: o boleto real do Itaú de referência é assim. Não
        #  se usa a data impressa na folha — ela seria uma data solta, sem a
        #  conferência que o DARF tem.
        self.assertIn("vencimento", app.motivo_fora_do_lancamento(
            self.doc(vencimento=None), CADASTRO))

    def test_darf_sem_vencimento_confirmado(self):
        d = app.documento_para_lancamento("g.pdf", dados_boleto=ARRECADACAO, texto="")
        self.assertIn("DARF", app.motivo_fora_do_lancamento(d, CADASTRO))

    def test_nota_sem_vencimento_proprio_pode_ser_lancada(self):
        #  A nota não tem vencimento próprio — ele vem do lote. Não é motivo
        #  de exclusão.
        d = app.documento_para_lancamento("n.pdf", dados_nfse=NFSE_SEM_RETENCAO)
        self.assertEqual(app.motivo_fora_do_lancamento(d, CADASTRO), "")


class TestLoteHomogeneo(unittest.TestCase):

    def test_um_tipo_so_passa(self):
        docs = [app.documento_para_lancamento("a.pdf", dados_boleto=BOLETO),
                app.documento_para_lancamento("b.pdf", dados_boleto=BOLETO)]
        self.assertEqual(app.conferir_lote_homogeneo(docs), "")

    def test_lote_vazio_passa(self):
        self.assertEqual(app.conferir_lote_homogeneo([]), "")

    def test_nota_no_meio_de_boletos_recusa_e_diz_qual(self):
        docs = [app.documento_para_lancamento("boleto.pdf", dados_boleto=BOLETO),
                app.documento_para_lancamento("nota_perdida.pdf", dados_nfse=NFSE_SEM_RETENCAO)]
        mensagem = app.conferir_lote_homogeneo(docs)
        self.assertTrue(mensagem)
        self.assertIn("nota_perdida.pdf", mensagem)

    def test_boleto_e_darf_sao_tipos_diferentes(self):
        docs = [app.documento_para_lancamento("b.pdf", dados_boleto=BOLETO),
                app.documento_para_lancamento("g.pdf", dados_boleto=ARRECADACAO, texto=DARF_OK)]
        self.assertTrue(app.conferir_lote_homogeneo(docs))


class TestMontarLancamentos(unittest.TestCase):

    def test_boleto_vira_lancamento_com_as_seis_colunas(self):
        docs = [app.documento_para_lancamento("b.pdf", dados_boleto=BOLETO)]
        lancamentos = app.montar_lancamentos(docs, CADASTRO)
        self.assertEqual(len(lancamentos), 1)
        l = lancamentos[0]
        self.assertEqual(set(l), set(app.COLUNAS_POR_DOCUMENTO))
        self.assertEqual(l["condominio"], "44")
        self.assertEqual(l["valor"], 82.9)
        self.assertEqual(l["vencimento"], datetime.date(2026, 9, 10))
        self.assertEqual(l["linha_digitavel"], BOLETO["linha_digitavel"])
        self.assertIsNone(l["numero_documento"])
        self.assertIsNone(l["competencia"])

    def test_nota_usa_o_vencimento_do_lote(self):
        docs = [app.documento_para_lancamento("n.pdf", dados_nfse=NFSE_SEM_RETENCAO)]
        l = app.montar_lancamentos(docs, CADASTRO, datetime.date(2026, 9, 5))[0]
        self.assertEqual(l["vencimento"], datetime.date(2026, 9, 5))
        self.assertEqual(l["numero_documento"], "12367")
        self.assertEqual(l["competencia"], datetime.date(2026, 8, 24))
        self.assertIsNone(l["linha_digitavel"])

    def test_nota_sem_vencimento_do_lote_e_erro_de_quem_chama(self):
        docs = [app.documento_para_lancamento("n.pdf", dados_nfse=NFSE_SEM_RETENCAO)]
        with self.assertRaises(ValueError):
            app.montar_lancamentos(docs, CADASTRO)

    def test_documentos_com_motivo_ficam_de_fora(self):
        docs = [app.documento_para_lancamento("ok.pdf", dados_boleto=BOLETO),
                app.documento_para_lancamento("lago.pdf", dados_boleto=dict(BOLETO, documento_pagador=LAGO)),
                app.documento_para_lancamento("sem_venc.pdf", dados_boleto=dict(BOLETO, vencimento=None))]
        lancamentos = app.montar_lancamentos(docs, CADASTRO)
        self.assertEqual(len(lancamentos), 1)

    def test_ordem_dos_documentos_e_preservada(self):
        SAN_REMO = "08578541000103"      # ID SL 42
        docs = [app.documento_para_lancamento("1.pdf", dados_boleto=dict(BOLETO, documento_pagador=SAN_REMO)),
                app.documento_para_lancamento("2.pdf", dados_boleto=BOLETO)]
        self.assertEqual([l["condominio"] for l in app.montar_lancamentos(docs, CADASTRO)],
                         ["42", "44"])


import shutil
import tempfile

from openpyxl import Workbook, load_workbook

#  Modelo sintético com as colunas que importam aqui. O molde traz lixo nas
#  colunas por documento DE PROPÓSITO: o programa tem de sobrescrevê-las,
#  nunca herdar.
CABECALHO = ["condomínio", "vencimento", "competencia", "fornecedor",
             "conta_categoria", "numero_documento", "complemento", "valor",
             "linha_digitavel", "chave"]
MOLDE = [None, None, None, "4521", "2.1.49 Envio Informações E-Social",
         "LIXO DO MOLDE", "Exames Médicos", None, "LIXO DO MOLDE", 46]


def montar_modelo(caminho):
    wb = Workbook()
    sheet = wb.active
    sheet.append(CABECALHO)
    sheet.append(MOLDE)
    wb.save(caminho)
    return caminho


class TestGerarComColunasPorDocumento(unittest.TestCase):

    def setUp(self):
        self.pasta = tempfile.mkdtemp()
        self.modelo = montar_modelo(os.path.join(self.pasta, "modelo.xlsx"))
        self.saida = os.path.join(self.pasta, "saida.xlsx")

    def tearDown(self):
        shutil.rmtree(self.pasta, ignore_errors=True)

    def gerar(self, lancamentos, **kwargs):
        app.gerar_planilha_despesas(self.modelo, self.saida, lancamentos, **kwargs)
        sheet = load_workbook(self.saida).active
        colunas = {app._normalizar_cabecalho(c.value): c.column for c in sheet[1]}
        return sheet, colunas

    def lancamento(self, **mudancas):
        base = {"condominio": "44", "valor": 82.9,
                "vencimento": datetime.date(2026, 9, 10),
                "linha_digitavel": "23790472089000015979956006290706515650000008290",
                "numero_documento": None, "competencia": None}
        base.update(mudancas)
        return base

    def test_escreve_as_colunas_por_documento(self):
        sheet, col = self.gerar([self.lancamento()])
        self.assertEqual(sheet.cell(2, col["condominio"]).value, "44")
        self.assertEqual(sheet.cell(2, col["valor"]).value, 82.9)
        self.assertEqual(sheet.cell(2, col["linha_digitavel"]).value,
                         "23790472089000015979956006290706515650000008290")

    def test_data_sai_como_data_de_verdade(self):
        sheet, col = self.gerar([self.lancamento()])
        celula = sheet.cell(2, col["vencimento"])
        self.assertTrue(celula.is_date)
        self.assertEqual(celula.number_format, app.FORMATO_DATA_DESPESAS)

    def test_none_apaga_o_que_o_molde_tinha(self):
        #  "—" no spec é célula vazia, mesmo que o molde tenha algo ali.
        sheet, col = self.gerar([self.lancamento(linha_digitavel=None)])
        self.assertIsNone(sheet.cell(2, col["linha_digitavel"]).value)
        self.assertIsNone(sheet.cell(2, col["numero_documento"]).value)

    def test_colunas_do_molde_continuam_sendo_copiadas(self):
        sheet, col = self.gerar([self.lancamento(), self.lancamento()])
        for linha in (2, 3):
            self.assertEqual(sheet.cell(linha, col["fornecedor"]).value, "4521")
            self.assertEqual(sheet.cell(linha, col["complemento"]).value, "Exames Médicos")
            self.assertEqual(sheet.cell(linha, col["chave"]).value, 46)

    def test_chave_do_lote_continua_valendo(self):
        sheet, col = self.gerar([self.lancamento()], chave=99)
        self.assertEqual(sheet.cell(2, col["chave"]).value, 99)

    def test_modelo_sem_uma_coluna_por_documento_avisa(self):
        wb = Workbook()
        wb.active.append(["condomínio", "valor"])
        wb.active.append([None, None])
        wb.save(self.modelo)
        with self.assertRaises(ValueError) as erro:
            self.gerar([self.lancamento()])
        self.assertIn("linha_digitavel", str(erro.exception))

    def test_formato_antigo_de_tupla_continua_funcionando(self):
        #  A aba 3 continua chamando com (id_sl, valor).
        sheet, col = self.gerar([("44", 57.75)])
        self.assertEqual(sheet.cell(2, col["condominio"]).value, "44")
        self.assertEqual(sheet.cell(2, col["valor"]).value, 57.75)


class TestListarModelosLancamento(unittest.TestCase):

    def setUp(self):
        self.pasta = tempfile.mkdtemp()

    def tearDown(self):
        shutil.rmtree(self.pasta, ignore_errors=True)

    def criar(self, nome):
        open(os.path.join(self.pasta, nome), "wb").close()

    def test_lista_os_xlsx_pelo_nome_sem_extensao_em_ordem(self):
        self.criar("Sindicato - Contr. Assistencial.xlsx")
        self.criar("FF - Exames Medicos.xlsx")
        rotulos = [r for r, _ in app.listar_modelos_lancamento(self.pasta)]
        self.assertEqual(rotulos, ["FF - Exames Medicos", "Sindicato - Contr. Assistencial"])

    def test_devolve_o_caminho_do_arquivo(self):
        self.criar("FF - PCMSO.xlsx")
        (_, caminho), = app.listar_modelos_lancamento(self.pasta)
        self.assertEqual(caminho, os.path.join(self.pasta, "FF - PCMSO.xlsx"))

    def test_ignora_o_arquivo_de_trava_do_excel(self):
        #  Com um modelo aberto no Excel aparece "~$Nome.xlsx" na pasta; não
        #  é modelo e não pode virar opção.
        self.criar("FF - PCMSO.xlsx")
        self.criar("~$FF - PCMSO.xlsx")
        self.assertEqual(len(app.listar_modelos_lancamento(self.pasta)), 1)

    def test_ignora_o_que_nao_e_xlsx(self):
        self.criar("leia-me.txt")
        self.criar("antigo.xls")
        self.assertEqual(app.listar_modelos_lancamento(self.pasta), [])

    def test_pasta_que_nao_existe(self):
        self.assertEqual(app.listar_modelos_lancamento(os.path.join(self.pasta, "nao")), [])


if __name__ == "__main__":
    unittest.main()
