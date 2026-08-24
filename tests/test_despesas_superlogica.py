import os
import shutil
import sys
import tempfile
import datetime
import unittest
from decimal import Decimal

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

_RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if _RAIZ not in sys.path:
    sys.path.insert(0, _RAIZ)
_AQUI = os.path.dirname(os.path.abspath(__file__))
if _AQUI not in sys.path:
    sys.path.insert(0, _AQUI)

import logica as app
from cadastro_teste import CADASTRO_TESTE

KLOSTERS = {"arquivo": "a.pdf", "codigo": "10004", "condominio": "10004 KLOSTERS",
            "valor": Decimal("57.75")}
SAN_REMO = {"arquivo": "b.pdf", "codigo": "10002", "condominio": "10002 SAN REMO",
            "valor": Decimal("11.55")}
#  Cadastrado, mas sem ID no Superlógica — 10 dos 764 reais estão assim.
LAGO = {"arquivo": "c.pdf", "codigo": "10590", "condominio": "10590 LAGO MAGGIORE",
        "valor": Decimal("7.70")}
#  Resolvido à mão depois de "não foi possível ler": não tem código nenhum.
SEM_CODIGO = {"arquivo": "d.pdf", "codigo": None, "condominio": "",
              "valor": Decimal("300.30")}


def resultado_com(processados=(), pendentes=(), ignorados=()):
    return {"processados": list(processados), "pendentes": list(pendentes),
            "ignorados": list(ignorados)}

#  Modelo sintético no formato do arquivo real do Superlógica: cabeçalho na
#  linha 1 e uma linha de exemplo na 2, com os campos que se repetem
#  preenchidos e condomínio/valor em branco. Nomes de coluna com acento e
#  caixa mista de propósito — é assim que vêm do Superlógica.
CABECALHO_MODELO = [
    "condomínio", "vencimento", "fornecedor", "conta_categoria",
    "valor", "forma_de_pagamento", "chave",
]
MOLDE_MODELO = [
    None, None, "DINAMICA SERVICOS POSTAIS", "2.4.6 Correios - Postagem Simples",
    None, "Trans. bancária", 46,
]


def montar_modelo(caminho, cabecalho=None, molde=None, linhas_extras=0):
    """Escreve um modelo sintético no disco e devolve o caminho."""
    wb = Workbook()
    sheet = wb.active
    sheet.append(cabecalho if cabecalho is not None else CABECALHO_MODELO)
    for celula in sheet[1]:
        celula.font = Font(bold=True)
    sheet.append(molde if molde is not None else MOLDE_MODELO)
    for _ in range(linhas_extras):
        sheet.append(molde if molde is not None else MOLDE_MODELO)
    wb.save(caminho)
    return caminho


class TestGerarPlanilhaDespesas(unittest.TestCase):
    def setUp(self):
        self.pasta = tempfile.mkdtemp()
        self.modelo = montar_modelo(os.path.join(self.pasta, "modelo.xlsx"))
        self.saida = os.path.join(self.pasta, "despesas.xlsx")

    def tearDown(self):
        shutil.rmtree(self.pasta, ignore_errors=True)

    def gerar(self, lancamentos):
        app.gerar_planilha_despesas(self.modelo, self.saida, lancamentos)
        return load_workbook(self.saida).active

    def test_uma_linha_por_lancamento_na_ordem_dada(self):
        sheet = self.gerar([("45", Decimal("7.70")),
                            ("496", Decimal("138.60")),
                            ("496", Decimal("119.35"))])
        #  Cabeçalho + 3 lançamentos, sem sobra
        self.assertEqual(sheet.max_row, 4)
        self.assertEqual([sheet.cell(row=l, column=1).value for l in (2, 3, 4)],
                         ["45", "496", "496"])
        self.assertEqual([sheet.cell(row=l, column=5).value for l in (2, 3, 4)],
                         [7.70, 138.60, 119.35])

    def test_campos_do_molde_repetem_em_todas_as_linhas(self):
        sheet = self.gerar([("45", Decimal("7.70")), ("48", Decimal("57.75"))])
        for linha in (2, 3):
            self.assertEqual(sheet.cell(row=linha, column=3).value,
                             "DINAMICA SERVICOS POSTAIS")
            self.assertEqual(sheet.cell(row=linha, column=4).value,
                             "2.4.6 Correios - Postagem Simples")
            self.assertEqual(sheet.cell(row=linha, column=6).value, "Trans. bancária")
            self.assertEqual(sheet.cell(row=linha, column=7).value, 46)

    def test_linha_molde_nao_sobra_em_branco(self):
        """A linha 2 do modelo é o molde e vira a primeira linha real — o
        arquivo final não pode ter nenhuma linha com condomínio/valor vazios."""
        sheet = self.gerar([("45", Decimal("7.70"))])
        self.assertEqual(sheet.max_row, 2)
        self.assertEqual(sheet.cell(row=2, column=1).value, "45")
        self.assertEqual(sheet.cell(row=2, column=5).value, 7.70)

    def test_linhas_extras_do_modelo_nao_sobram(self):
        """Modelo salvo com várias linhas de exemplo não pode deixar resto."""
        modelo = montar_modelo(os.path.join(self.pasta, "modelo3.xlsx"),
                               linhas_extras=4)
        app.gerar_planilha_despesas(modelo, self.saida, [("45", Decimal("7.70"))])
        sheet = load_workbook(self.saida).active
        self.assertEqual(sheet.max_row, 2)

    def test_valor_gravado_como_numero_nao_texto(self):
        sheet = self.gerar([("45", Decimal("7.70"))])
        self.assertIsInstance(sheet.cell(row=2, column=5).value, float)

    def test_acha_as_colunas_por_nome_nao_por_posicao(self):
        """O layout é do Superlógica e pode mudar: as colunas são localizadas
        pelo nome normalizado, com acento e caixa ignorados."""
        modelo = montar_modelo(
            os.path.join(self.pasta, "outra_ordem.xlsx"),
            cabecalho=["VALOR", "fornecedor", "  Condomínio  "],
            molde=[None, "DINAMICA SERVICOS POSTAIS", None])
        app.gerar_planilha_despesas(modelo, self.saida, [("45", Decimal("7.70"))])
        sheet = load_workbook(self.saida).active
        self.assertEqual(sheet.cell(row=2, column=1).value, 7.70)
        self.assertEqual(sheet.cell(row=2, column=3).value, "45")

    def test_modelo_sem_coluna_condominio_avisa(self):
        modelo = montar_modelo(os.path.join(self.pasta, "sem_cond.xlsx"),
                               cabecalho=["valor", "fornecedor"],
                               molde=[None, "DINAMICA"])
        with self.assertRaises(ValueError) as erro:
            app.gerar_planilha_despesas(modelo, self.saida, [("45", Decimal("7.70"))])
        self.assertIn("condominio", str(erro.exception))

    def test_modelo_sem_coluna_valor_avisa(self):
        modelo = montar_modelo(os.path.join(self.pasta, "sem_valor.xlsx"),
                               cabecalho=["condomínio", "fornecedor"],
                               molde=[None, "DINAMICA"])
        with self.assertRaises(ValueError) as erro:
            app.gerar_planilha_despesas(modelo, self.saida, [("45", Decimal("7.70"))])
        self.assertIn("valor", str(erro.exception))

    def test_lista_vazia_nao_gera_arquivo_pela_metade(self):
        with self.assertRaises(ValueError):
            app.gerar_planilha_despesas(self.modelo, self.saida, [])
        self.assertFalse(os.path.isfile(self.saida))

    def test_estilo_do_molde_e_preservado_nas_linhas_novas(self):
        """Copiar em vez de reconstruir é o ponto do desenho: o importador do
        Superlógica pode depender de formato de célula."""
        wb = load_workbook(self.modelo)
        wb.active.cell(row=2, column=3).font = Font(bold=True, italic=True)
        wb.save(self.modelo)
        sheet = self.gerar([("45", Decimal("7.70")), ("48", Decimal("57.75"))])
        for linha in (2, 3):
            self.assertTrue(sheet.cell(row=linha, column=3).font.bold)
            self.assertTrue(sheet.cell(row=linha, column=3).font.italic)


class TestColunasDeData(unittest.TestCase):
    """O Excel guarda data como número de série e só o formato da célula diz
    que aquilo é data. Um modelo com a célula em "General" fazia o Superlógica
    ler o número cru e gravar 01/01/1970 — aconteceu de verdade, e as linhas
    foram recusadas na importação."""

    def setUp(self):
        self.pasta = tempfile.mkdtemp()
        self.saida = os.path.join(self.pasta, "despesas.xlsx")

    def tearDown(self):
        shutil.rmtree(self.pasta, ignore_errors=True)

    def modelo_com_vencimento(self, valor, formato=None, nome="vencimento"):
        caminho = os.path.join(self.pasta, "m.xlsx")
        wb = Workbook()
        sheet = wb.active
        sheet.append(["condomínio", nome, "fornecedor", "valor"])
        sheet.append([None, valor, "DINAMICA", None])
        if formato:
            sheet.cell(row=2, column=2).number_format = formato
        wb.save(caminho)
        return caminho

    def test_numero_de_serie_vira_data_de_verdade(self):
        #  46255 é 21/08/2026 no calendário do Excel
        modelo = self.modelo_com_vencimento(46255)
        app.gerar_planilha_despesas(modelo, self.saida,
                                    [("45", Decimal("7.70")), ("48", Decimal("57.75"))])
        sheet = load_workbook(self.saida).active
        for linha in (2, 3):
            celula = sheet.cell(row=linha, column=2)
            self.assertTrue(celula.is_date, f"linha {linha} não saiu como data")
            self.assertEqual(celula.value, datetime.datetime(2026, 8, 21))

    def test_data_de_verdade_no_modelo_e_preservada(self):
        modelo = self.modelo_com_vencimento(datetime.datetime(2026, 8, 21),
                                            formato="DD/MM/YYYY")
        app.gerar_planilha_despesas(modelo, self.saida, [("45", Decimal("7.70"))])
        celula = load_workbook(self.saida).active.cell(row=2, column=2)
        self.assertTrue(celula.is_date)
        self.assertEqual(celula.value, datetime.datetime(2026, 8, 21))

    def test_vencimento_vazio_continua_vazio(self):
        """O modelo original não preenche vencimento — isso é legítimo."""
        modelo = self.modelo_com_vencimento(None)
        app.gerar_planilha_despesas(modelo, self.saida, [("45", Decimal("7.70"))])
        self.assertIsNone(load_workbook(self.saida).active.cell(row=2, column=2).value)

    def test_texto_no_lugar_da_data_e_recusado(self):
        modelo = self.modelo_com_vencimento("21/08/2026")
        with self.assertRaises(ValueError) as erro:
            app.gerar_planilha_despesas(modelo, self.saida, [("45", Decimal("7.70"))])
        self.assertIn("vencimento", str(erro.exception))

    def test_numero_fora_da_faixa_de_datas_e_recusado(self):
        """123 como série seria 1900 — não é vencimento de ninguém, então é
        mais provável que alguém tenha digitado outra coisa na coluna."""
        modelo = self.modelo_com_vencimento(123)
        with self.assertRaises(ValueError) as erro:
            app.gerar_planilha_despesas(modelo, self.saida, [("45", Decimal("7.70"))])
        self.assertIn("vencimento", str(erro.exception))

    def test_competencia_recebe_o_mesmo_tratamento(self):
        modelo = self.modelo_com_vencimento(46255, nome="competencia")
        app.gerar_planilha_despesas(modelo, self.saida, [("45", Decimal("7.70"))])
        celula = load_workbook(self.saida).active.cell(row=2, column=2)
        self.assertTrue(celula.is_date)

    def test_liquidacao_com_acento_no_cabecalho_tambem(self):
        modelo = self.modelo_com_vencimento(46255, nome="liquidação")
        app.gerar_planilha_despesas(modelo, self.saida, [("45", Decimal("7.70"))])
        celula = load_workbook(self.saida).active.cell(row=2, column=2)
        self.assertTrue(celula.is_date)


class TestConverterDataDigitada(unittest.TestCase):
    """O vencimento é dado do lote, não do modelo — o programa pergunta a cada
    geração, como já faz com a tarifa. Antes ele era digitado na célula do
    modelo, e foi ali que a data virou número e a importação quebrou."""

    def test_formato_brasileiro(self):
        data, aviso = app.converter_data_digitada("21/08/2026")
        self.assertEqual(data, datetime.datetime(2026, 8, 21))
        self.assertEqual(aviso, "")

    def test_aceita_traco_e_ponto_como_separador(self):
        for texto in ("21-08-2026", "21.08.2026"):
            self.assertEqual(app.converter_data_digitada(texto)[0],
                             datetime.datetime(2026, 8, 21), texto)

    def test_aceita_ano_de_dois_digitos(self):
        self.assertEqual(app.converter_data_digitada("21/08/26")[0],
                         datetime.datetime(2026, 8, 21))

    def test_aceita_espacos_nas_pontas(self):
        self.assertEqual(app.converter_data_digitada("  21/08/2026 ")[0],
                         datetime.datetime(2026, 8, 21))

    def test_recusa_data_que_nao_existe(self):
        """31/02 não é erro de digitação inocente — recusar é melhor que
        deslizar para 03/03."""
        data, aviso = app.converter_data_digitada("31/02/2026")
        self.assertIsNone(data)
        self.assertTrue(aviso)

    def test_recusa_formato_americano_ambiguo(self):
        """2026-08-21 não é o formato que se digita aqui; aceitar convidaria a
        confundir dia com mês em datas como 03/04."""
        data, _ = app.converter_data_digitada("2026-08-21")
        self.assertIsNone(data)

    def test_recusa_texto_e_vazio(self):
        for texto in ("amanhã", "", "   ", "21/08"):
            data, aviso = app.converter_data_digitada(texto)
            self.assertIsNone(data, texto)
            self.assertTrue(aviso, texto)


class TestVencimentoNaGeracao(unittest.TestCase):
    def setUp(self):
        self.pasta = tempfile.mkdtemp()
        self.modelo = montar_modelo(os.path.join(self.pasta, "modelo.xlsx"))
        self.saida = os.path.join(self.pasta, "despesas.xlsx")

    def tearDown(self):
        shutil.rmtree(self.pasta, ignore_errors=True)

    def test_vencimento_informado_preenche_todas_as_linhas_como_data(self):
        app.gerar_planilha_despesas(
            self.modelo, self.saida,
            [("45", Decimal("7.70")), ("48", Decimal("57.75"))],
            vencimento=datetime.datetime(2026, 8, 21))
        sheet = load_workbook(self.saida).active
        for linha in (2, 3):
            celula = sheet.cell(row=linha, column=2)
            self.assertEqual(celula.value, datetime.datetime(2026, 8, 21))
            self.assertTrue(celula.is_date, f"linha {linha} não saiu como data")

    def test_vencimento_informado_vence_o_que_esta_no_modelo(self):
        """Se alguém tiver deixado uma data velha no modelo, a do lote manda."""
        wb = load_workbook(self.modelo)
        wb.active.cell(row=2, column=2).value = datetime.datetime(2020, 1, 1)
        wb.active.cell(row=2, column=2).number_format = "DD/MM/YYYY"
        wb.save(self.modelo)
        app.gerar_planilha_despesas(self.modelo, self.saida,
                                    [("45", Decimal("7.70"))],
                                    vencimento=datetime.datetime(2026, 8, 21))
        celula = load_workbook(self.saida).active.cell(row=2, column=2)
        self.assertEqual(celula.value, datetime.datetime(2026, 8, 21))

    def test_sem_vencimento_o_modelo_continua_mandando(self):
        """Chamada sem vencimento é o comportamento de antes — o modelo decide."""
        app.gerar_planilha_despesas(self.modelo, self.saida,
                                    [("45", Decimal("7.70"))])
        self.assertIsNone(load_workbook(self.saida).active.cell(row=2, column=2).value)

    def test_modelo_sem_coluna_vencimento_avisa_ao_informar_data(self):
        modelo = montar_modelo(os.path.join(self.pasta, "sem_venc.xlsx"),
                               cabecalho=["condomínio", "fornecedor", "valor"],
                               molde=[None, "DINAMICA", None])
        with self.assertRaises(ValueError) as erro:
            app.gerar_planilha_despesas(modelo, self.saida,
                                        [("45", Decimal("7.70"))],
                                        vencimento=datetime.datetime(2026, 8, 21))
        self.assertIn("vencimento", str(erro.exception))


class TestNormalizarCabecalho(unittest.TestCase):
    def test_tira_acento_caixa_e_espacos(self):
        self.assertEqual(app._normalizar_cabecalho("  Condomínio "), "condominio")
        self.assertEqual(app._normalizar_cabecalho("VALOR"), "valor")

    def test_celula_vazia_vira_string_vazia(self):
        self.assertEqual(app._normalizar_cabecalho(None), "")


class TestLancamentosDeDespesa(unittest.TestCase):
    def test_um_lancamento_por_protocolo_na_ordem(self):
        lancamentos, travas = app.lancamentos_de_despesa(
            resultado_com([KLOSTERS, SAN_REMO]), CADASTRO_TESTE)
        self.assertEqual(lancamentos, [("44", Decimal("57.75")),
                                       ("42", Decimal("11.55"))])
        self.assertEqual(travas, {"sem_valor": [], "sem_id_sl": []})

    def test_mesmo_condominio_duas_vezes_gera_dois_lancamentos(self):
        """Cada protocolo físico vira um lançamento — somar quebraria a
        correspondência com o papel que originou cada um."""
        outro = dict(KLOSTERS, arquivo="a2.pdf", valor=Decimal("138.60"))
        lancamentos, _ = app.lancamentos_de_despesa(
            resultado_com([KLOSTERS, outro]), CADASTRO_TESTE)
        self.assertEqual(lancamentos, [("44", Decimal("57.75")),
                                       ("44", Decimal("138.60"))])

    def test_pendente_sem_valor_trava(self):
        lancamentos, travas = app.lancamentos_de_despesa(
            resultado_com([KLOSTERS], pendentes=[{"arquivo": "p.pdf"}]),
            CADASTRO_TESTE)
        self.assertEqual(travas["sem_valor"], ["p.pdf"])
        self.assertEqual(travas["sem_id_sl"], [])

    def test_condominio_sem_id_sl_trava(self):
        lancamentos, travas = app.lancamentos_de_despesa(
            resultado_com([KLOSTERS, LAGO]), CADASTRO_TESTE)
        self.assertEqual(travas["sem_id_sl"], ["c.pdf"])
        self.assertEqual(travas["sem_valor"], [])

    def test_protocolo_sem_codigo_trava_por_falta_de_id(self):
        _, travas = app.lancamentos_de_despesa(
            resultado_com([SEM_CODIGO]), CADASTRO_TESTE)
        self.assertEqual(travas["sem_id_sl"], ["d.pdf"])

    def test_codigo_fora_do_cadastro_trava_por_falta_de_id(self):
        fora = dict(KLOSTERS, arquivo="e.pdf", codigo="99999")
        _, travas = app.lancamentos_de_despesa(
            resultado_com([fora]), CADASTRO_TESTE)
        self.assertEqual(travas["sem_id_sl"], ["e.pdf"])

    def test_arquivo_que_nao_e_protocolo_nao_trava_nem_vira_lancamento(self):
        """Ignorados nunca deveriam virar despesa — não podem travar o lote."""
        lancamentos, travas = app.lancamentos_de_despesa(
            resultado_com([KLOSTERS], ignorados=[{"arquivo": "outro.pdf"}]),
            CADASTRO_TESTE)
        self.assertEqual(lancamentos, [("44", Decimal("57.75"))])
        self.assertEqual(travas, {"sem_valor": [], "sem_id_sl": []})

    def test_lote_inteiro_travado_devolve_lista_vazia(self):
        lancamentos, travas = app.lancamentos_de_despesa(
            resultado_com([LAGO], pendentes=[{"arquivo": "p.pdf"}]),
            CADASTRO_TESTE)
        self.assertEqual(lancamentos, [])
        self.assertEqual(travas["sem_id_sl"], ["c.pdf"])
        self.assertEqual(travas["sem_valor"], ["p.pdf"])

    def test_valor_continua_decimal(self):
        """Dinheiro só vira float na escrita da planilha."""
        lancamentos, _ = app.lancamentos_de_despesa(
            resultado_com([KLOSTERS]), CADASTRO_TESTE)
        self.assertIsInstance(lancamentos[0][1], Decimal)


if __name__ == "__main__":
    unittest.main()


class TestBuscarCondominios(unittest.TestCase):
    """Busca do seletor de condomínio do painel: o funcionário digita parte do
    nome ou do código e escolhe na lista. Existe porque um protocolo pode
    chegar sem código legível (manuscrito por cima) ou com código fora do
    cadastro — e nesses casos nada mais no painel destravava a geração."""

    def test_acha_por_pedaco_do_nome(self):
        achados = app.buscar_condominios("klost", CADASTRO_TESTE)
        self.assertEqual([c["codigo"] for c in achados], ["10004"])

    def test_busca_ignora_acento_e_caixa(self):
        achados = app.buscar_condominios("BONFIM", CADASTRO_TESTE)
        self.assertEqual(len(achados), 2)

    def test_acha_por_codigo(self):
        achados = app.buscar_condominios("10002", CADASTRO_TESTE)
        self.assertEqual([c["nome"] for c in achados], ["SAN REMO"])

    def test_traz_o_id_sl_junto(self):
        """É o que o painel precisa para destravar a geração."""
        achados = app.buscar_condominios("klost", CADASTRO_TESTE)
        self.assertEqual(achados[0]["id_sl"], "44")

    def test_termo_vazio_traz_tudo_ordenado_por_nome(self):
        achados = app.buscar_condominios("", CADASTRO_TESTE)
        self.assertEqual(len(achados), len(CADASTRO_TESTE))
        nomes = [c["nome"] for c in achados]
        self.assertEqual(nomes, sorted(nomes))

    def test_sem_correspondencia_devolve_lista_vazia(self):
        self.assertEqual(app.buscar_condominios("zzzz", CADASTRO_TESTE), [])

    def test_traz_tambem_quem_esta_sem_id_sl(self):
        """Esconder seria pior: a pessoa escolheria e nada aconteceria, sem
        entender por quê. Melhor aparecer e o painel avisar que falta o ID."""
        achados = app.buscar_condominios("maggiore", CADASTRO_TESTE)
        self.assertEqual(achados[0]["id_sl"], "")
