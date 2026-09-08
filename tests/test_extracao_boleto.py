"""
Leitura do código de barras / linha digitável dos boletos (aba 2).

As duas linhas digitáveis usadas aqui vieram de boletos reais (CAIXA e Itaú)
e estão INTEIRAS de propósito: o que se está testando é justamente a
conferência dos dígitos verificadores, e um número inventado à mão não
provaria nada — passaria ou falharia por construção. As fixtures em volta
delas, sim, são fabricadas, com um condomínio do cadastro de teste.
"""

import datetime
import os
import sys
import tempfile
import unittest

_RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if _RAIZ not in sys.path:
    sys.path.insert(0, _RAIZ)
_AQUI = os.path.dirname(os.path.abspath(__file__))
if _AQUI not in sys.path:
    sys.path.insert(0, _AQUI)

import logica as app
from cadastro_teste import CADASTRO_TESTE

KLOSTERS = "01195716000154"

#  CAIXA, vencimento 10/09/2026, R$ 250,00.
LINHA_CAIXA = "10498057553218010004500000763854115650000025000"
#  Itaú, R$ 188,48, fator 9999 — este emissor não põe o vencimento na barra,
#  embora a folha traga 15/09/2026 impresso.
LINHA_ITAU = "34191090087507589929640762040000899990000018848"


def ler(nome):
    with open(os.path.join(_AQUI, "dados", nome), encoding="utf-8") as f:
        return f.read()


class TestDigitosVerificadores(unittest.TestCase):
    def test_linhas_reais_sao_validas(self):
        self.assertTrue(app.linha_digitavel_valida(LINHA_CAIXA))
        self.assertTrue(app.linha_digitavel_valida(LINHA_ITAU))

    def test_aceita_com_pontuacao_como_vem_impresso(self):
        self.assertTrue(app.linha_digitavel_valida(
            "10498.05755  32180.100045  00000.763854  1  15650000025000"))

    def test_um_digito_trocado_derruba_a_linha(self):
        # É esta recusa que autoriza confiar em valor e vencimento lidos daqui:
        # dígito errado não vira valor errado na planilha, vira linha recusada.
        for posicao in (0, 12, 30, 40, 46):
            trocado = LINHA_CAIXA[:posicao] + str(
                (int(LINHA_CAIXA[posicao]) + 1) % 10) + LINHA_CAIXA[posicao + 1:]
            self.assertFalse(app.linha_digitavel_valida(trocado),
                             f"dígito {posicao} trocado passou na conferência")

    def test_comprimento_errado_nao_e_linha_digitavel(self):
        self.assertFalse(app.linha_digitavel_valida(LINHA_CAIXA[:-1]))
        self.assertFalse(app.linha_digitavel_valida(LINHA_CAIXA + "0"))
        self.assertFalse(app.linha_digitavel_valida(""))


class TestConversaoBarra(unittest.TestCase):
    def test_linha_vira_codigo_de_barras_e_volta(self):
        codigo = app.linha_digitavel_para_codigo_barras(LINHA_CAIXA)
        self.assertEqual(len(codigo), 44)
        self.assertEqual(codigo[:3], "104")
        self.assertTrue(app.codigo_barras_valido(codigo))
        self.assertEqual(app.codigo_barras_para_linha_digitavel(codigo), LINHA_CAIXA)

    def test_formatacao_igual_a_impressa_no_papel(self):
        self.assertEqual(
            app.formatar_linha_digitavel(LINHA_CAIXA),
            "10498.05755 32180.100045 00000.763854 1 15650000025000")


class TestFatorDeVencimento(unittest.TestCase):
    def test_fator_do_boleto_real_da_caixa(self):
        # 1565, já no ciclo reiniciado em 22/02/2025 — bate com a data
        # impressa na folha.
        self.assertEqual(app.vencimento_do_fator("1565"), datetime.date(2026, 9, 10))

    def test_primeiro_e_ultimo_fator_do_ciclo_em_vigor(self):
        self.assertEqual(app.vencimento_do_fator("1000"), datetime.date(2025, 2, 22))
        self.assertEqual(app.vencimento_do_fator("9998"), datetime.date(2049, 10, 12))

    def test_fator_abaixo_de_1000_nao_vira_data(self):
        # Faixa só usada pelo ciclo antigo, cujas datas são todas passado.
        self.assertIsNone(app.vencimento_do_fator("0999"))

    def test_sentinelas_nao_viram_data(self):
        self.assertIsNone(app.vencimento_do_fator("0000"))
        self.assertIsNone(app.vencimento_do_fator("9999"))


class TestArrecadacao(unittest.TestCase):
    """Contas de concessionária (água, luz, gás) usam o código de arrecadação:
    48 dígitos, começando com 8, sem banco e sem fator de vencimento."""

    def _montar(self, codigo44):
        return app.codigo_barras_para_linha_digitavel(codigo44)

    def test_ida_e_volta_com_modulo_10(self):
        # identificador de valor 6 -> DVs em módulo 10
        codigo = ("8260" + "00000012345" + "0" * 29)[:44]
        linha = self._montar(codigo)
        self.assertEqual(len(linha), 48)
        self.assertEqual(app.linha_digitavel_para_codigo_barras(linha), codigo)

    def test_valor_e_segmento_lidos_do_codigo(self):
        #  8=arrecadação, 2=saneamento, 6=valor efetivo em módulo 10,
        #  0=DV geral (só o valor importa para esta leitura), depois 11
        #  dígitos de valor: 123,45.
        codigo = ("8260" + "00000012345" + "0" * 29)[:44]
        dados = app.dados_do_codigo_barras(codigo)
        self.assertEqual(dados["tipo"], "Arrecadação")
        self.assertIn("SANEAMENTO", dados["emissor"])
        self.assertEqual(dados["valor"], 123.45)
        self.assertIsNone(dados["vencimento"])

    def test_arrecadacao_nunca_tem_47_digitos(self):
        self.assertFalse(app.linha_digitavel_valida("8" + "0" * 46))


class TestExtracaoDoTexto(unittest.TestCase):
    def test_acha_a_linha_no_meio_do_texto_do_boleto(self):
        texto = ler("boleto_caixa.txt")
        self.assertEqual(app.extrair_linha_digitavel(texto), LINHA_CAIXA)

    def test_numeros_soltos_da_pagina_nao_viram_linha_digitavel(self):
        # CNPJ, datas, nosso número, CEP: nenhum passa na conferência.
        texto = ("CNPJ 11.222.333/0001-81 03/09/2026 13711 22050-020\n"
                 "14/180000000007638-1 6357/0805753-2 250,00\n")
        self.assertIsNone(app.extrair_linha_digitavel(texto))

    def test_espaco_a_mais_no_meio_do_numero_nao_derruba_a_leitura(self):
        # Saída real do pypdf num boleto do Itaú: o extrator enfia um espaço
        # dentro do último campo e cola o "341-7" que vem depois, então o
        # trecho tem 50 dígitos em vez de 47.
        texto = ("LITROS POR SEGUNDO MANUTENCAO PREDIAL \n"
                 "34191.09008 75075.899296 40762.040000 8 99990000018 848 341-7 \n")
        self.assertEqual(app.extrair_linha_digitavel(texto), LINHA_ITAU)

    def test_linha_partida_em_duas_linhas_de_texto(self):
        texto = "VALOR\n34191.09008 75075.899296 40762.040000\n8 99990000018848\n341-7\n"
        self.assertEqual(app.extrair_linha_digitavel(texto), LINHA_ITAU)

    def test_documento_que_nao_e_boleto_devolve_none(self):
        self.assertIsNone(app.extrair_dados_boleto(ler("nfse_ff.txt"), CADASTRO_TESTE))

    def test_dados_do_boleto_da_caixa(self):
        dados = app.extrair_dados_boleto(ler("boleto_caixa.txt"), CADASTRO_TESTE)
        self.assertEqual(dados["tipo"], "Boleto bancário")
        self.assertEqual(dados["emissor"], "104 - CAIXA ECONÔMICA FEDERAL")
        self.assertEqual(dados["vencimento"], datetime.date(2026, 9, 10))
        self.assertEqual(dados["valor"], 250.0)
        self.assertEqual(dados["documento_pagador"], KLOSTERS)

    def test_pagador_e_o_cadastrado_nunca_o_beneficiario(self):
        # Os dois CNPJs aparecem na folha; só o condomínio está no cadastro.
        dados = app.extrair_dados_boleto(ler("boleto_caixa.txt"), CADASTRO_TESTE)
        self.assertNotEqual(dados["documento_pagador"], "11222333000181")

    def test_sem_nenhum_cnpj_cadastrado_ninguem_e_escolhido(self):
        dados = app.extrair_dados_boleto(ler("boleto_sem_vencimento.txt"), CADASTRO_TESTE)
        self.assertIsNone(dados["documento_pagador"])
        self.assertEqual(dados["valor"], 188.48)
        self.assertIsNone(dados["vencimento"])


class TestLinhaDaPlanilha(unittest.TestCase):
    def setUp(self):
        self.dados = app.extrair_dados_boleto(ler("boleto_caixa.txt"), CADASTRO_TESTE)

    def test_colunas_da_planilha(self):
        # Linha digitável, barra, condomínio (nome e código), vencimento e
        # valor — o que a planilha existe para mostrar; arquivo e observação
        # são o entorno.
        linha = app.linha_planilha_boleto("boleto.pdf", self.dados, CADASTRO_TESTE)
        self.assertEqual(linha[0], "boleto.pdf")
        self.assertEqual(linha[1], app.formatar_linha_digitavel(LINHA_CAIXA))
        self.assertEqual(linha[2], app.linha_digitavel_para_codigo_barras(LINHA_CAIXA))
        self.assertEqual(linha[3], "KLOSTERS")
        self.assertEqual(linha[4], "10004")
        self.assertEqual(linha[5], datetime.date(2026, 9, 10))
        self.assertEqual(linha[6], 250.0)
        self.assertEqual(linha[7], "")

    def test_linha_digitavel_igual_a_impressa_no_boleto(self):
        # É o número que a pessoa lê na parte de cima do papel: sai pontuado,
        # como impresso, para conferir a olho e digitar no banco.
        linha = app.linha_planilha_boleto("boleto.pdf", self.dados, CADASTRO_TESTE)
        self.assertEqual(linha[1],
                         "10498.05755 32180.100045 00000.763854 1 15650000025000")

    def test_os_dois_numeros_saem_como_texto(self):
        # 47 e 44 dígitos: como número, o Excel viraria notação científica e
        # comeria os dígitos verificadores.
        linha = app.linha_planilha_boleto("boleto.pdf", self.dados, CADASTRO_TESTE)
        self.assertIsInstance(linha[1], str)
        self.assertIsInstance(linha[2], str)
        self.assertEqual(len(linha[2]), 44)

    def test_sem_cnpj_no_boleto_cai_para_o_nome_do_arquivo_e_avisa(self):
        dados = app.extrair_dados_boleto(ler("boleto_sem_vencimento.txt"), CADASTRO_TESTE)
        linha = app.linha_planilha_boleto("Boleto SAN REMO 09.26.pdf", dados,
                                          CADASTRO_TESTE)
        self.assertEqual(linha[3], "SAN REMO")
        self.assertEqual(linha[4], "10002")
        self.assertIn("nome do arquivo", linha[7])

    def test_condominio_nao_identificado_e_dito_na_observacao(self):
        dados = app.extrair_dados_boleto(ler("boleto_sem_vencimento.txt"), CADASTRO_TESTE)
        linha = app.linha_planilha_boleto("desconhecido.pdf", dados, CADASTRO_TESTE)
        self.assertEqual(linha[3], "")
        self.assertEqual(linha[4], "")
        self.assertIn("não identificado", linha[7])

    def test_sem_vencimento_na_barra_a_celula_vem_vazia_com_motivo(self):
        dados = app.extrair_dados_boleto(ler("boleto_sem_vencimento.txt"), CADASTRO_TESTE)
        linha = app.linha_planilha_boleto("desconhecido.pdf", dados, CADASTRO_TESTE)
        self.assertIsNone(linha[5])
        self.assertIn("sem vencimento", linha[7])

    def test_arquivo_ignorado_mantem_o_formato_da_linha(self):
        linha = app.linha_planilha_boleto("x.pdf", None, CADASTRO_TESTE, "Erro ao ler")
        self.assertEqual(len(linha), len(app.COLUNAS_BOLETO))
        self.assertEqual(linha[-1], "Erro ao ler")


class TestPlanilhaComDuasAbas(unittest.TestCase):
    def setUp(self):
        self.pasta = tempfile.mkdtemp()
        self.caminho = os.path.join(self.pasta, "extracao.xlsx")
        dados = app.extrair_dados_boleto(ler("boleto_caixa.txt"), CADASTRO_TESTE)
        self.linha_boleto = app.linha_planilha_boleto("boleto.pdf", dados, CADASTRO_TESTE)

    def test_aba_de_boletos_so_existe_quando_ha_boleto(self):
        from openpyxl import load_workbook
        app.salvar_planilha_nfse(self.caminho, [])
        self.assertEqual(load_workbook(self.caminho).sheetnames, ["Notas fiscais"])

        app.salvar_planilha_nfse(self.caminho, [], [self.linha_boleto])
        wb = load_workbook(self.caminho)
        self.assertEqual(wb.sheetnames, ["Notas fiscais", "Boletos"])
        self.assertEqual(
            [c.value for c in wb["Boletos"][1]],
            ["Arquivo", "Linha digitável", "Código de barras", "Condomínio",
             "Código", "Vencimento", "Valor", "Observação"])
        self.assertEqual(wb["Boletos"].cell(row=2, column=5).value, "10004")

    def test_valor_e_data_de_verdade_e_a_barra_como_texto(self):
        from openpyxl import load_workbook
        app.salvar_planilha_nfse(self.caminho, [], [self.linha_boleto])
        aba = load_workbook(self.caminho)["Boletos"]
        self.assertEqual(aba.cell(row=2, column=6).value,
                         datetime.datetime(2026, 9, 10))
        self.assertEqual(aba.cell(row=2, column=7).value, 250.0)
        self.assertEqual(aba.cell(row=2, column=2).value,
                         app.formatar_linha_digitavel(LINHA_CAIXA))
        self.assertEqual(aba.cell(row=2, column=3).value,
                         app.linha_digitavel_para_codigo_barras(LINHA_CAIXA))


if __name__ == "__main__":
    unittest.main()
