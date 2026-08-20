import os
import shutil
import sys
import tempfile
import unittest

_RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if _RAIZ not in sys.path:
    sys.path.insert(0, _RAIZ)
_AQUI = os.path.dirname(os.path.abspath(__file__))
if _AQUI not in sys.path:
    sys.path.insert(0, _AQUI)

import logica as app
from openpyxl import Workbook, load_workbook

KLOSTERS = "01195716000154"


class TestIdSl(unittest.TestCase):
    """
    "ID SL" é o código do condomínio no Superlógica — outro número, sem
    relação com o código interno (ex: KLOSTERS é 10004 aqui e 44 lá). Fica
    guardado no cadastro só como referência; nada da identificação nem dos
    carimbos usa esse campo.
    """

    def setUp(self):
        self.pasta = tempfile.mkdtemp()
        self.caminho = os.path.join(self.pasta, "cadastro.xlsx")

    def tearDown(self):
        shutil.rmtree(self.pasta, ignore_errors=True)

    def _planilha(self, cabecalho, linhas):
        wb = Workbook()
        sheet = wb.active
        sheet.append(cabecalho)
        for linha in linhas:
            sheet.append(linha)
        wb.save(self.caminho)

    def test_le_id_sl_da_quarta_coluna(self):
        self._planilha(["CNPJ", "Código", "Nome do Condomínio", "ID SL"],
                       [["01.195.716/0001-54", "10004", "KLOSTERS", "44"]])
        cadastro = app.carregar_cadastro(self.caminho)
        self.assertEqual(cadastro[KLOSTERS]["id_sl"], "44")

    def test_planilha_antiga_de_tres_colunas_continua_carregando(self):
        # Cadastro gravado por uma versão anterior do app não tem a coluna —
        # tem de carregar sem quebrar, com o campo vazio.
        self._planilha(["CNPJ", "Código", "Nome do Condomínio"],
                       [["01.195.716/0001-54", "10004", "KLOSTERS"]])
        cadastro = app.carregar_cadastro(self.caminho)
        self.assertEqual(cadastro[KLOSTERS]["id_sl"], "")
        self.assertEqual(cadastro[KLOSTERS]["codigo"], "10004")

    def test_id_sl_vazio_na_planilha_vira_string_vazia(self):
        self._planilha(["CNPJ", "Código", "Nome do Condomínio", "ID SL"],
                       [["01.195.716/0001-54", "10004", "KLOSTERS", None]])
        cadastro = app.carregar_cadastro(self.caminho)
        self.assertEqual(cadastro[KLOSTERS]["id_sl"], "")

    def test_salvar_preserva_o_id_sl(self):
        # O risco real: salvar_cadastro reescreve a planilha inteira. Se não
        # gravasse a coluna, uma edição qualquer na aba de Cadastro apagaria
        # o ID SL de todos os 747 condomínios de uma vez, em silêncio.
        app.salvar_cadastro(self.caminho, {
            KLOSTERS: {"codigo": "10004", "nome": "KLOSTERS", "id_sl": "44"},
        })
        cadastro = app.carregar_cadastro(self.caminho)
        self.assertEqual(cadastro[KLOSTERS]["id_sl"], "44")

    def test_salvar_registro_sem_id_sl_nao_quebra(self):
        # Registro criado pelo formulário antigo (sem a chave) tem de gravar
        # normalmente, com a coluna em branco.
        app.salvar_cadastro(self.caminho, {
            KLOSTERS: {"codigo": "10004", "nome": "KLOSTERS"},
        })
        cadastro = app.carregar_cadastro(self.caminho)
        self.assertEqual(cadastro[KLOSTERS]["id_sl"], "")

    def test_cabecalho_gravado_tem_a_coluna(self):
        app.salvar_cadastro(self.caminho, {
            KLOSTERS: {"codigo": "10004", "nome": "KLOSTERS", "id_sl": "44"},
        })
        sheet = load_workbook(self.caminho).active
        self.assertEqual([c.value for c in sheet[1]],
                         ["CNPJ", "Código", "Nome do Condomínio", "ID SL"])


if __name__ == "__main__":
    unittest.main()
