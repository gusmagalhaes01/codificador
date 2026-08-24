import os
import shutil
import sys
import tempfile
import unittest

_RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if _RAIZ not in sys.path:
    sys.path.insert(0, _RAIZ)
_AQUI = os.path.dirname(os.path.abspath(__file__))
if _AQUI not in sys.path:
    sys.path.insert(0, _AQUI)

import logica as app
from openpyxl import Workbook, load_workbook

CPF_VILA_MARINA = "52998224725"
KLOSTERS = "01195716000154"


class TestCadastroPorCpf(unittest.TestCase):
    def setUp(self):
        self.pasta = tempfile.mkdtemp()
        self.caminho = os.path.join(self.pasta, "cadastro.xlsx")

    def tearDown(self):
        shutil.rmtree(self.pasta, ignore_errors=True)

    def test_ida_e_volta_preserva_o_cpf(self):
        original = {
            CPF_VILA_MARINA: {"codigo": "11300", "nome": "VILA MARINA", "id_sl": "930"},
            KLOSTERS: {"codigo": "10004", "nome": "KLOSTERS", "id_sl": "44"},
        }
        app.salvar_cadastro(self.caminho, original)
        self.assertEqual(app.carregar_cadastro(self.caminho), original)

    def test_cpf_gravado_formatado_como_cpf(self):
        app.salvar_cadastro(self.caminho, {
            CPF_VILA_MARINA: {"codigo": "11300", "nome": "VILA MARINA", "id_sl": "930"},
        })
        sheet = load_workbook(self.caminho).active
        self.assertEqual(sheet.cell(row=2, column=1).value, "529.982.247-25")

    def test_cabecalho_menciona_os_dois_documentos(self):
        app.salvar_cadastro(self.caminho, {})
        sheet = load_workbook(self.caminho).active
        self.assertEqual(sheet.cell(row=1, column=1).value, "CNPJ / CPF")

    def test_planilha_antiga_so_com_cnpj_continua_carregando(self):
        """O cabeçalho mudou de nome, mas carregar_cadastro lê por posição —
        planilha gravada por uma versão anterior precisa continuar abrindo."""
        wb = Workbook()
        sheet = wb.active
        sheet.append(["CNPJ", "Código", "Nome do Condomínio", "ID SL"])
        sheet.append(["01.195.716/0001-54", "10004", "KLOSTERS", "44"])
        wb.save(self.caminho)
        cadastro = app.carregar_cadastro(self.caminho)
        self.assertEqual(cadastro[KLOSTERS]["nome"], "KLOSTERS")

    def test_cpf_e_cnpj_nunca_colidem_como_chave(self):
        cadastro = {
            CPF_VILA_MARINA: {"codigo": "11300", "nome": "VILA MARINA", "id_sl": "930"},
            KLOSTERS: {"codigo": "10004", "nome": "KLOSTERS", "id_sl": "44"},
        }
        app.salvar_cadastro(self.caminho, cadastro)
        self.assertEqual(len(app.carregar_cadastro(self.caminho)), 2)


if __name__ == "__main__":
    unittest.main()
